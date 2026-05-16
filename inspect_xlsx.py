import openpyxl
import os

directory = r"D:\BSIT\3RD YR\Web Software Tools\Mini-project\Talais_miniproject\sf 9 card"
files = [f for f in os.listdir(directory) if f.endswith(".xlsx")]

for filename in files:
    path = os.path.join(directory, filename)
    print(f"\n{'='*20}\nFILE: {filename}")
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        print(f"Sheets: {wb.sheetnames}")
        
        for sheetname in wb.sheetnames:
            ws = wb[sheetname]
            print(f"\n  SHEET: {sheetname}")
            print(f"  Used Range: {ws.dimensions}")
            
            # Fix for MultiCellRange length issue
            merged_count = 0
            try:
                merged_count = len(list(ws.merged_cells.ranges))
            except:
                pass
            print(f"  Merged Cells Count: {merged_count}")
            
            cells_data = []
            non_empty_count = 0
            
            for row in range(1, 41):
                for col in range(1, 16):
                    cell = ws.cell(row=row, column=col)
                    val = cell.value
                    coord = cell.coordinate
                    if val is not None:
                        if non_empty_count < 30:
                            cells_data.append(f"{coord}: {val}")
                        non_empty_count += 1
                
            print(f"  First 30 non-empty cells:")
            for item in cells_data[:30]:
                print(f"    {item}")
                
            # Identifying grade-level or subject labels
            labels = []
            for row in ws.iter_rows(min_row=1, max_row=20, min_col=1, max_col=10):
                for cell in row:
                    if cell.value and isinstance(cell.value, str):
                        val_lower = cell.value.lower()
                        if "grade" in val_lower or "subject" in val_lower or "learning area" in val_lower:
                            labels.append(f"{cell.coordinate}: {cell.value}")
            if labels:
                print(f"  Potential Labels:")
                for l in labels:
                    print(f"    {l}")
                    
    except Exception as e:
        print(f"  Error processing file: {e}")
